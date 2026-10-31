"""Offline DOCX exports keep source styling independent of bilingual order."""

from __future__ import annotations

import hashlib

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from wenyi_core.assemble.writer import assemble
from wenyi_core.config import Config
from wenyi_core.ingest.docx_reader import read_docx
from wenyi_core.pipeline.runstore import ExportSnapshotStore


@pytest.fixture
def bilingual_docx(tmp_path):
    source = tmp_path / "source.docx"
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    for container, mixed in [
        (doc, False),
        (doc, True),
        (table.cell(0, 0), False),
        (table.cell(0, 1), True),
    ]:
        paragraph = container.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run("Red " if mixed else "Uniform source.")
        run.bold = True
        run.font.size = Pt(14)
        run.font.color.rgb = RGBColor.from_string("112233")
        if mixed:
            paragraph.add_run("plain")
        else:
            run.italic = True
            run.underline = True
    doc.save(source)

    book = read_docx(str(source), "en", "zh")
    chapter = book.chapters[0]
    for segment in chapter.segments:
        segment.target = "普通红色" if segment.source == "Red plain" else "统一样式译文。"
        styles = segment.meta.get("docx_styles")
        if styles:
            styles["target_digest"] = hashlib.sha256(segment.target.encode("utf-8")).hexdigest()
            styles["placements"] = [
                {
                    "target_start": 2,
                    "target_end": 4,
                    "bold": True,
                    "color": "112233",
                    "size_pt": 14,
                }
            ]
    snapshot = ExportSnapshotStore(
        str(tmp_path / "run"),
        {
            "fmt": "docx",
            "source_lang": "en",
            "target_lang": "zh",
            "chapters": [{"index": chapter.index}],
        },
        {chapter.index: chapter},
    )
    return source, snapshot


@pytest.mark.parametrize("order", ["target_first", "source_first"])
@pytest.mark.parametrize("preserve_source_style", [False, True])
def test_bilingual_docx_source_style_is_readable_and_order_independent(
    tmp_path, bilingual_docx, order, preserve_source_style
):
    source, snapshot = bilingual_docx
    before = snapshot.load_chapter(0).model_dump()
    cfg = Config.from_dict(
        {
            "output": {
                "bilingual": True,
                "bilingual_order": order,
                "bilingual_preserve_source_style": preserve_source_style,
            }
        }
    )
    written = assemble(
        snapshot,
        str(source),
        out_path=str(tmp_path / "bilingual.docx"),
        out_format="docx",
        bilingual=cfg.output.bilingual,
        order=cfg.output.bilingual_order,
        preserve_source_style=cfg.output.bilingual_preserve_source_style,
    )
    result = Document(written)
    assert len(result.paragraphs) == 4
    assert len(result.tables) == 1
    sides = []
    for index in (0, 2):
        first, second = result.paragraphs[index : index + 2]
        assert first.alignment == second.alignment == WD_ALIGN_PARAGRAPH.CENTER
        sides.append((first.runs, second.runs))
    for cell in result.tables[0].rows[0].cells:
        assert len(cell.paragraphs) == 1
        paragraph = cell.paragraphs[0]
        separator = next(i for i, run in enumerate(paragraph.runs) if run.text == "\n")
        sides.append((paragraph.runs[:separator], paragraph.runs[separator + 1 :]))

    for index, (first, second) in enumerate(sides):
        source_runs, target_runs = (first, second) if order == "source_first" else (second, first)
        mixed = index % 2 == 1
        assert "".join(run.text for run in source_runs) == (
            "Red plain" if mixed else "Uniform source."
        )
        assert "".join(run.text for run in target_runs) == (
            "普通红色" if mixed else "统一样式译文。"
        )
        for run in source_runs:
            assert run.font.highlight_color is None
            assert run.font.name != "宋体"
        if preserve_source_style:
            assert source_runs[0].text == ("Red " if mixed else "Uniform source.")
            assert source_runs[0].bold
            assert str(source_runs[0].font.color.rgb) == "112233"
            assert source_runs[0].font.size == Pt(14)
            if mixed:
                assert source_runs[1].text == "plain"
                assert source_runs[1].bold is None
                assert source_runs[1].font.color.rgb is None
            else:
                assert source_runs[0].italic
                assert source_runs[0].underline
        else:
            for run in source_runs:
                assert run.font.color.rgb is None
                assert run.font.size is None
                assert run.bold is None
                assert run.italic is None
                assert run.underline is None
        styled_target = target_runs[-1] if mixed else target_runs[0]
        assert styled_target.text == ("红色" if mixed else "统一样式译文。")
        assert styled_target.bold
        assert str(styled_target.font.color.rgb) == "112233"
        assert styled_target.font.size == Pt(14)
        assert all(run.font.name == "宋体" for run in target_runs)
    assert snapshot.load_chapter(0).model_dump() == before


@pytest.mark.parametrize("order", ["target_first", "source_first"])
@pytest.mark.parametrize("preserve_source_style", [False, True])
@pytest.mark.parametrize("target", [None, "", "Uniform source."])
def test_bilingual_docx_source_fallback_is_not_duplicated_or_restyled(
    tmp_path, bilingual_docx, order, preserve_source_style, target
):
    source, snapshot = bilingual_docx
    chapter = snapshot.load_chapter(0)
    for segment in chapter.segments:
        segment.source = "Uniform source."
        segment.target = target
        segment.meta.pop("docx_styles", None)
        segment.meta["docx_style"] = {"bold": True, "color": "112233", "size_pt": 14}
    snapshot = ExportSnapshotStore(
        str(tmp_path / "fallback"), snapshot.load_manifest(), {chapter.index: chapter}
    )
    written = assemble(
        snapshot,
        str(source),
        out_path=str(tmp_path / "fallback.docx"),
        out_format="docx",
        bilingual=True,
        order=order,
        preserve_source_style=preserve_source_style,
    )
    result = Document(written)
    paragraphs = result.paragraphs + [cell.paragraphs[0] for cell in result.tables[0].rows[0].cells]
    assert len(paragraphs) == 4
    for paragraph in paragraphs:
        assert paragraph.text == "Uniform source."
        assert paragraph.runs[0].bold
        assert str(paragraph.runs[0].font.color.rgb) == "112233"
        assert paragraph.runs[0].font.size == Pt(14)
        assert paragraph.runs[0].font.highlight_color is None
        assert paragraph.runs[0].font.name != "宋体"
